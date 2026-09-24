import io
import datetime
import pandas as pd
import numpy as np
from typing import Dict, Any, List, Optional, Tuple
from sqlalchemy.orm import Session
from rapidfuzz import process, fuzz
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.models import EmployeeMaster, OrgUnitAlias

def clean_str(val: Any) -> Optional[str]:
    if pd.isna(val) or val is None:
        return None
    s = str(val).strip()
    return s if s and s != '-' and s.lower() != 'nan' else None

def clean_nik(val: Any) -> Optional[str]:
    if pd.isna(val) or val is None:
        return None
    s = str(val).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s if s and s != '-' and s.lower() != 'nan' else None

class MasterDataPipeline:
    def __init__(self, db: Optional[Session] = None, fuzzy_threshold: float = 85.0):
        self.db = db
        self.fuzzy_threshold = fuzzy_threshold
        self.alias_map: Dict[str, str] = {}
        self.canonical_units: Dict[str, List[str]] = {
            "direktorat": [],
            "kompartemen": [],
            "departemen": []
        }
        self.load_aliases()

    def load_aliases(self):
        if self.db:
            aliases = self.db.query(OrgUnitAlias).all()
            for a in aliases:
                self.alias_map[a.nama_asli.lower().strip()] = a.nama_kanonik
                if a.nama_kanonik not in self.canonical_units.get(a.jenis, []):
                    self.canonical_units.setdefault(a.jenis, []).append(a.nama_kanonik)

    def normalize_unit_name(self, name: Optional[str], unit_type: str, unknown_list: List[str]) -> Optional[str]:
        if not name:
            return None
        name_clean = str(name).strip()
        name_lower = name_clean.lower()

        # 1. Exact lookup in alias map
        if name_lower in self.alias_map:
            return self.alias_map[name_lower]

        # 2. Check canonical list
        canonicals = self.canonical_units.get(unit_type, [])
        for c in canonicals:
            if c.lower() == name_lower:
                return c

        # 3. Fuzzy match if canonical names exist
        if canonicals:
            match = process.extractOne(name_clean, canonicals, scorer=fuzz.token_sort_ratio)
            if match and match[1] >= self.fuzzy_threshold:
                return match[0]

        # 4. Built-in standard abbreviation expansion
        expanded = name_clean
        if expanded.startswith("Komp. "):
            expanded = "Kompartemen " + expanded[6:]
        elif expanded.startswith("Komp "):
            expanded = "Kompartemen " + expanded[5:]
        elif expanded.startswith("Dep. "):
            expanded = "Departemen " + expanded[5:]
        elif expanded.startswith("Dep "):
            expanded = "Departemen " + expanded[4:]
        elif expanded.startswith("Bag. "):
            expanded = "Bagian " + expanded[5:]
        elif expanded.startswith("Bag "):
            expanded = "Bagian " + expanded[4:]

        if expanded.lower() in self.alias_map:
            return self.alias_map[expanded.lower()]

        if expanded != name_clean and expanded not in canonicals:
            unknown_list.append(f"{unit_type}: {name_clean} (expanded to {expanded})")
        elif name_clean not in canonicals and name_clean not in self.alias_map.values():
            unknown_list.append(f"{unit_type}: {name_clean}")

        return expanded

    def classify_employee(self, row: pd.Series, snapshot_date: datetime.date) -> str:
        # If row explicitly comes from dedicated non-active sheet (e.g. PKWT sheet, Pensiun sheet, PI sheet)
        forced = str(row.get('_forced_kategori', '') or '').strip().upper()
        if forced in ['PURNA', 'PKWT', 'PI']:
            return forced

        statuspegawai = str(row.get('statuspegawai', '') or row.get('status_pegawai', '') or '').strip().upper()
        status_kerja = str(row.get('status', '') or row.get('status_kerja', '') or '').strip().upper()
        kompartemen = str(row.get('kompartemen', '') or '').strip().upper()
        departemen = str(row.get('departemen', '') or '').strip().upper()
        postitle = str(row.get('postitle', '') or row.get('pos_title', '') or row.get('nama_posisi', '') or row.get('position_title', '') or '').strip().upper()
        nm_jabatan = str(row.get('nm_jabatan', '') or '').strip().upper()
        regu = str(row.get('regu', '') or '').strip().upper()
        tgl_pen = row.get('tgl_pen', None)
        # 1. PURNA / Pensiun
        # Check tgl_pen <= snapshot_date (e.g. 2026-09-22)
        if pd.notna(tgl_pen):
            try:
                if isinstance(tgl_pen, (datetime.date, datetime.datetime)):
                    pen_date = tgl_pen.date() if isinstance(tgl_pen, datetime.datetime) else tgl_pen
                    if pen_date <= snapshot_date and str(pen_date) not in ['1900-01-01', '0001-01-01']:
                        return "PURNA"
                else:
                    tgl_str = str(tgl_pen).strip()
                    if tgl_str not in ['', '-', '1900-01-01', 'NaT', 'nan', 'None']:
                        p_dt = pd.to_datetime(tgl_pen, errors='coerce')
                        if pd.notna(p_dt) and p_dt.date() <= snapshot_date and str(p_dt.date()) not in ['1900-01-01', '0001-01-01']:
                            return "PURNA"
            except Exception:
                pass

        if (
            'PURNA' in statuspegawai or
            'PENSIUN' in statuspegawai
        ):
            return "PURNA"

        # 2. Penugasan PT Pupuk Indonesia (PI) - Check kompartemen/departemen contains 'KARYAWAN PENEMPATAN' or 'KARYAWAN PENUGASAN' or PI postitle/regu
        if (
            'KARYAWAN PENEMPATAN' in kompartemen or
            'KARYAWAN PENEMPATAN' in departemen or
            'PENEMPATAN' in kompartemen or
            'PENEMPATAN' in departemen or
            'KARYAWAN PENUGASAN' in kompartemen or
            'KARYAWAN PENUGASAN' in departemen or
            'PT PUPUK INDONESIA' in regu or
            'PENUGASAN PT PUPUK INDONESIA' in postitle or
            'PENEMPATAN PT PUPUK INDONESIA' in postitle
        ):
            return "PI"

        # 3. PKWT / Pro Hire (Active)
        if any(k in statuspegawai for k in ['PKWT', 'PRO_HIRE', 'PRO HIRE', 'KONTRAK']):
            return "PKWT"

        # 4. Karyawan Aktif
        return "AKTIF"

    def process_workbook(self, excel_file: io.BytesIO, snapshot_date: Optional[datetime.date] = None) -> Dict[str, Any]:
        if snapshot_date is None:
            snapshot_date = datetime.date.today()

        xls = pd.ExcelFile(excel_file)
        sheet_names = xls.sheet_names

        # Detect 'Tidak Coaching' / 'Belum Coaching' sheet
        tidak_coaching_sheet = next(
            (s for s in sheet_names if ('tidak' in s.lower() or 'belum' in s.lower() or 'tdk' in s.lower()) and 'coaching' in s.lower()),
            None
        )
        tidak_coaching_niks = set()
        tidak_coaching_names = set()
        df_tc = None
        if tidak_coaching_sheet:
            try:
                df_tc = pd.read_excel(xls, tidak_coaching_sheet)
                df_tc.columns = [str(c).strip().lower().replace(' ', '_') for c in df_tc.columns]
                nik_col = next((c for c in df_tc.columns if 'nik' in c or 'badge' in c or 'nip' in c or 'no_peg' in c), None)
                if nik_col:
                    for n in df_tc[nik_col].dropna():
                        c_n = clean_nik(n)
                        if c_n:
                            tidak_coaching_niks.add(c_n)
                nama_col = next((c for c in df_tc.columns if 'nama' in c), None)
                if nama_col:
                    for n in df_tc[nama_col].dropna():
                        c_n = clean_str(n)
                        if c_n:
                            tidak_coaching_names.add(c_n.lower().strip())
            except Exception as e:
                print("Error reading tidak coaching sheet:", e)

        # Case 1: Workbook contains specific curated tabs
        if 'Karyawan Aktif (DOP)' in sheet_names or 'Karyawan Aktif' in sheet_names:
            akt_sheet = 'Karyawan Aktif (DOP)' if 'Karyawan Aktif (DOP)' in sheet_names else 'Karyawan Aktif'
            df_akt = pd.read_excel(xls, akt_sheet)
            dfs = [df_akt]

            if 'PKWT' in sheet_names:
                df_pkwt = pd.read_excel(xls, 'PKWT')
                df_pkwt['_forced_kategori'] = 'PKWT'
                dfs.append(df_pkwt)

            if 'Karyawan Pensiun' in sheet_names:
                df_pen = pd.read_excel(xls, 'Karyawan Pensiun')
                df_pen['_forced_kategori'] = 'PURNA'
                dfs.append(df_pen)

            if 'Penugasan PI' in sheet_names or 'Karyawan PI' in sheet_names:
                sheet_pi = 'Penugasan PI' if 'Penugasan PI' in sheet_names else 'Karyawan PI'
                df_pi = pd.read_excel(xls, sheet_pi)
                df_pi['_forced_kategori'] = 'PI'
                dfs.append(df_pi)

            combined_df = pd.concat(dfs, ignore_index=True)
            return self.process_dataframe(combined_df, snapshot_date, tidak_coaching_niks=tidak_coaching_niks, tidak_coaching_names=tidak_coaching_names)

        # Case 2: Single sheet or un-split raw data
        df = pd.read_excel(xls, sheet_name=0)
        return self.process_dataframe(df, snapshot_date, tidak_coaching_niks=tidak_coaching_niks, tidak_coaching_names=tidak_coaching_names)

    def process_dataframe(
        self,
        df: pd.DataFrame,
        snapshot_date: Optional[datetime.date] = None,
        tidak_coaching_niks: Optional[set] = None,
        tidak_coaching_names: Optional[set] = None
    ) -> Dict[str, Any]:
        if snapshot_date is None:
            snapshot_date = datetime.date.today()
        if tidak_coaching_niks is None:
            tidak_coaching_niks = set()
        if tidak_coaching_names is None:
            tidak_coaching_names = set()

        # Normalize column headers
        df.columns = [str(c).strip().lower().replace(' ', '_') for c in df.columns]

        unknown_units: List[str] = []
        processed_rows = []

        for idx, row in df.iterrows():
            nik_val = clean_nik(row.get('nik', None))
            nik_sap_val = clean_nik(row.get('nik_sap', None))
            nama_val = clean_str(row.get('nama', None)) or "Unknown"

            if not nik_val and not nik_sap_val:
                continue

            if not nik_val and nik_sap_val:
                nik_val = nik_sap_val

            # Normalize organization names
            direktorat = self.normalize_unit_name(clean_str(row.get('direktorat')), 'direktorat', unknown_units)
            kompartemen = self.normalize_unit_name(clean_str(row.get('kompartemen')), 'kompartemen', unknown_units)
            departemen = self.normalize_unit_name(clean_str(row.get('departemen')), 'departemen', unknown_units)
            bagian = clean_str(row.get('bagian'))
            seksi = clean_str(row.get('seksi'))
            regu = clean_str(row.get('regu'))
            poscode = clean_str(row.get('poscode')) or clean_str(row.get('pos_code')) or clean_str(row.get('position_code')) or clean_str(row.get('kode_posisi')) or clean_str(row.get('id_posisi')) or clean_str(row.get('pos_id'))
            postitle = clean_str(row.get('postitle')) or clean_str(row.get('pos_title')) or clean_str(row.get('position_title')) or clean_str(row.get('nama_posisi')) or clean_str(row.get('nm_jabatan'))
            eselon = clean_str(row.get('eselon'))
            nm_jabatan = clean_str(row.get('nm_jabatan'))
            status_pegawai = clean_str(row.get('statuspegawai'))
            status_kerja = clean_str(row.get('status')) or clean_str(row.get('status_kerja'))

            # Classification
            kategori = self.classify_employee(row, snapshot_date)
            if kategori == "PURNA":
                status_pegawai = "PURNA BAKTI"
            elif kategori == "PI":
                status_pegawai = "PENUGASAN PI"

            # Master file does NOT set delegasi; delegasi is strictly determined when KPI file is uploaded
            is_delegasi = False
            delegasi_posisi_lain = None

            # Tidak Coaching flag from tab (check by NIK or by Nama)
            is_tidak_coaching = False
            cleaned_nama_lower = nama_val.lower().strip() if nama_val else None
            if (
                nik_val in tidak_coaching_niks
                or (nik_sap_val and nik_sap_val in tidak_coaching_niks)
                or (cleaned_nama_lower and cleaned_nama_lower in tidak_coaching_names)
            ):
                is_tidak_coaching = True

            processed_rows.append({
                "nik": nik_val,
                "nik_sap": nik_sap_val,
                "nama": nama_val,
                "eselon": eselon,
                "nm_jabatan": nm_jabatan,
                "direktorat": direktorat,
                "kompartemen": kompartemen,
                "departemen": departemen,
                "bagian": bagian,
                "seksi": seksi,
                "regu": regu,
                "poscode": poscode,
                "postitle": postitle,
                "status_pegawai": status_pegawai,
                "status_kerja": status_kerja,
                "tgl_pen": clean_str(row.get('tgl_pen')),
                "kategori": kategori,
                "is_delegasi": is_delegasi,
                "delegasi_posisi_lain": delegasi_posisi_lain,
                "is_tidak_coaching": is_tidak_coaching,
                "snapshot_date": snapshot_date,
            })

        processed_df = pd.DataFrame(processed_rows)

        # Count per category
        kategori_counts = processed_df['kategori'].value_counts().to_dict() if not processed_df.empty else {}
        aktif_count = int(kategori_counts.get("AKTIF", 0))
        pkwt_count = int(kategori_counts.get("PKWT", 0))
        purna_count = int(kategori_counts.get("PURNA", 0))
        pi_count = int(kategori_counts.get("PI", 0))
        delegasi_count = int(processed_df['is_delegasi'].sum()) if not processed_df.empty else 0
        tidak_coaching_count = int(processed_df['is_tidak_coaching'].sum()) if not processed_df.empty else 0

        # Calculate Rekap (Direktorat & Kompartemen breakdown)
        rekap = []
        if not processed_df.empty:
            grouped = processed_df.groupby(['direktorat', 'kompartemen', 'kategori']).size().unstack(fill_value=0)
            for (dir_name, komp_name), row_data in grouped.iterrows():
                rekap.append({
                    "direktorat": dir_name,
                    "kompartemen": komp_name,
                    "aktif": int(row_data.get("AKTIF", 0)),
                    "pkwt": int(row_data.get("PKWT", 0)),
                    "purna": int(row_data.get("PURNA", 0)),
                    "pi": int(row_data.get("PI", 0)),
                    "total": int(row_data.sum())
                })

        return {
            "records": processed_rows,
            "total_rows": len(processed_rows),
            "aktif_count": aktif_count,
            "pkwt_count": pkwt_count,
            "purna_count": purna_count,
            "pi_count": pi_count,
            "delegasi_count": delegasi_count,
            "tidak_coaching_count": tidak_coaching_count,
            "unknown_units_count": len(set(unknown_units)),
            "unknown_units": list(set(unknown_units)),
            "rekap": rekap,
            "snapshot_date": snapshot_date,
        }

    def save_to_db(self, records: List[Dict[str, Any]], snapshot_date: datetime.date, source: str = "OVERVIEW") -> int:
        if not self.db:
            return 0

        # Replace previous master data snapshot for this source entirely so numbers are always fresh
        self.db.query(EmployeeMaster).filter(EmployeeMaster.source == source).delete()
        for r in records:
            r['source'] = source
        db_objects = [EmployeeMaster(**r) for r in records]
        self.db.bulk_save_objects(db_objects)
        self.db.commit()
        return len(db_objects)

    def generate_cleaned_master_excel(self, records: List[Dict[str, Any]], scope: str = "kompartemen") -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data Master Karyawan"

        # Query Planning and Coaching data if DB session available
        planning_map = {}
        coaching_map = {}
        if self.db:
            from app.models.models import PerformancePlanning, PerformanceCoaching
            for p in self.db.query(PerformancePlanning).all():
                nik_c = clean_nik(p.employee_nik)
                if nik_c and nik_c not in planning_map:
                    planning_map[nik_c] = p
            for c in self.db.query(PerformanceCoaching).all():
                nik_c = clean_nik(c.employee_nik)
                if nik_c and nik_c not in coaching_map:
                    coaching_map[nik_c] = c

        # Exact 19 requested columns
        headers = [
            "nik_sap",
            "nama",
            "eselon",
            "statuspegawai",
            "nm_jabatan",
            "direktorat",
            "kompartemen",
            "departemen",
            "bagian",
            "seksi",
            "regu",
            "poscode",
            "postitle",
            "Re Planning KPI",
            "Submitted Date",
            "Approved At",
            "Hasil Re Planning KPI",
            "Coaching Triwulan II",
            "Hasil Coaching"
        ]

        header_fill = PatternFill(start_color="0A5C36", end_color="0A5C36", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=10)
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        for c_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c_idx, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for r_idx, r in enumerate(records, 2):
            nik = clean_nik(r.get('nik')) or clean_nik(r.get('nik_sap'))
            nik_sap = clean_nik(r.get('nik_sap')) or clean_nik(r.get('nik')) or '-'
            nama = r.get('nama') or '-'
            eselon = r.get('eselon') or '-'
            statuspegawai = r.get('status_pegawai') or r.get('statuspegawai') or '-'
            nm_jabatan = r.get('nm_jabatan') or '-'
            direktorat = r.get('direktorat') or '-'
            kompartemen = r.get('kompartemen') or '-'
            departemen = r.get('departemen') or '-'
            bagian = r.get('bagian') or '-'
            seksi = r.get('seksi') or '-'
            regu = r.get('regu') or '-'
            poscode = r.get('poscode') or '-'
            postitle = r.get('postitle') or r.get('nm_jabatan') or '-'

            # Re Planning KPI & Coaching TW II
            plan_obj = planning_map.get(nik)
            coach_obj = coaching_map.get(nik)

            if plan_obj:
                re_planning_kpi = plan_obj.status_individu or "Belum"
                sub_date = str(plan_obj.submitted_date) if plan_obj.submitted_date else "-"
                app_date = str(plan_obj.approved_date) if plan_obj.approved_date else "-"
                status_ind_lower = (plan_obj.status_individu or '').lower()
                if 'approved' in status_ind_lower:
                    hasil_planning = "Approved"
                elif 'wait' in status_ind_lower or 'submit' in status_ind_lower:
                    hasil_planning = "Waiting Approval"
                elif 'draft' in status_ind_lower:
                    hasil_planning = "Drafted"
                else:
                    hasil_planning = plan_obj.status_individu or "Belum Planning"
            else:
                re_planning_kpi = "Belum"
                sub_date = "-"
                app_date = "-"
                hasil_planning = "Belum Planning"

            if coach_obj:
                c_status = (coach_obj.status or '').strip()
                c_status_lower = c_status.lower()
                if c_status_lower in ['approved', 'sudah', 'done']:
                    coaching_tw2 = "Sudah"
                else:
                    coaching_tw2 = "Belum"
                hasil_coaching = f"{c_status} ({coach_obj.jumlah_coaching or 0} Sesi)" if c_status else "Belum Coaching"
            else:
                coaching_tw2 = "Belum"
                hasil_coaching = "Belum Coaching"

            row_values = [
                nik_sap,
                nama,
                eselon,
                statuspegawai,
                nm_jabatan,
                direktorat,
                kompartemen,
                departemen,
                bagian,
                seksi,
                regu,
                poscode,
                postitle,
                re_planning_kpi,
                sub_date,
                app_date,
                hasil_planning,
                coaching_tw2,
                hasil_coaching
            ]

            for c_idx, val in enumerate(row_values, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=str(val) if val is not None else '-')
                cell.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(c.value or '')) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 3, 12), 35)

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

